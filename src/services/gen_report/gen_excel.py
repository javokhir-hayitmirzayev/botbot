from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from src.utilities.types import QUESTION_TYPES

# --- CONSTANTS FOR STYLES ---
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4F81BD") # Blue
CORRECT_FILL = PatternFill("solid", fgColor="C6EFCE") # Green
WRONG_FILL = PatternFill("solid", fgColor="FFC7CE") # Red
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
WRAP_ALIGN = Alignment(wrap_text=True, vertical="center")
THIN_BORDER = Border(left=Side(style='thin'), right=Side(style='thin'), 
                     top=Side(style='thin'), bottom=Side(style='thin'))

def create_excel_structure(ws):
    """Sets up headers and column widths."""
    headers = ["Question", "Multi-Correct", "A", "B", "C", "D", "E", "Correct Answer", "Your Answer", "Is Correct"]
    ws.append(headers)

    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 12
    for col in ['C', 'D', 'E', 'F', 'G']:
        ws.column_dimensions[col].width = 25
    ws.column_dimensions['H'].width = 15
    ws.column_dimensions['I'].width = 15
    ws.column_dimensions['J'].width = 12
    
    ws.freeze_panes = 'A2'

def write_questions_to_excel(ws, questions):
    """Iterates through processed questions and writes them to the sheet."""
    row_num = 2
    correct_count = 0
    label_to_index = {'A': 0, 'B': 1, 'C': 2, 'D': 3, 'E': 4}

    for q in questions:
        option_texts = [""] * 5
        for opt in q['options']:
            idx = label_to_index.get(opt['label'].upper())
            if idx is not None and idx < 5:
                option_texts[idx] = opt['text']

        q['correct_labels'].sort()
        q['user_labels'].sort()
        
        is_correct = (q['correct_labels'] == q['user_labels']) and len(q['user_labels']) > 0
        if is_correct: 
            correct_count += 1

        row_data = [
            q['text'],
            "Yes" if q['type'] == QUESTION_TYPES.MULTI_CORRECT.value else "No",
            *option_texts,
            ", ".join(q['correct_labels']),
            ", ".join(q['user_labels']) if q['user_labels'] else "Skipped",
            "Yes" if is_correct else "No"
        ]
        ws.append(row_data)

        ws.cell(row=row_num, column=1).alignment = WRAP_ALIGN
        
        for c in range(3, 8):
            ws.cell(row=row_num, column=c).alignment = WRAP_ALIGN
            
        for c in [2, 8, 9, 10]:
            ws.cell(row=row_num, column=c).alignment = CENTER_ALIGN

        status_cell = ws.cell(row=row_num, column=10)
        if is_correct:
            status_cell.fill = CORRECT_FILL
            status_cell.font = Font(color="006100")
        else:
            status_cell.fill = WRONG_FILL
            status_cell.font = Font(color="9C0006")

        row_num += 1
        
    return row_num, correct_count

def write_summary_section(ws, start_row, metadata, total_q, correct_q):
    """Writes the scorecard at the bottom."""
    row_num = start_row + 2
    
    duration_str = "N/A"
    if metadata['finished_at'] and metadata['started_at']:
        duration_str = str(metadata['finished_at'] - metadata['started_at']).split('.')[0]

    summary_data = [
        ("OVERALL RESULTS", ""),
        ("Test Name", metadata['test_name']),
        ("Date", metadata['started_at'].strftime("%Y-%m-%d %H:%M")),
        ("Time Taken", duration_str),
        ("Total Questions", total_q),
        ("Correct Answers", correct_q),
        ("Incorrect/Skipped", total_q - correct_q),
        ("Final Score", f"{metadata['score']}%"),
    ]

    for i, (metric, value) in enumerate(summary_data):
        curr_row = row_num + i
        
        c1 = ws.cell(row=curr_row, column=1, value=metric)
        c1.font = Font(bold=True, size=14 if i == 0 else 11)
        
        if i > 0:
            c2 = ws.cell(row=curr_row, column=2, value=value)
            c2.alignment = Alignment(horizontal="left")
